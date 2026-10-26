from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from helpers import ProjectCase, job, make_resume

from jobfinder.agents.resume import ResumeAgent, ResumeError, replace_text_nodes, skill_items
from jobfinder.common import digest
from jobfinder.resume_layout import ResumeDocument, create_profile


class ResumeTests(ProjectCase):
    def agent(self):
        return ResumeAgent(self.load(), None)

    def target(self, agent, section="summary"):
        return next(index for index, entry in agent.targets.items() if entry["section"] == section)

    def test_shifted_document_generates_new_map(self):
        original = create_profile(self.settings.base_resume, self.settings.data["resume"])
        make_resume(self.settings.base_resume, shifted=True)
        shifted = create_profile(self.settings.base_resume, self.settings.data["resume"])
        self.assertNotEqual(
            original["targets"][0]["paragraph_id"], shifted["targets"][0]["paragraph_id"]
        )
        self.assertNotEqual(original["base_sha256"], shifted["base_sha256"])

    def test_identity_role_dates_and_education_protected(self):
        agent = self.agent()
        for i, paragraph in enumerate(agent.document.paragraphs):
            if (
                "Example Candidate" in paragraph.text
                or "@" in paragraph.text
                or "2021" in paragraph.text
                or "Degree" in paragraph.text
            ):
                self.assertNotIn(i, agent.targets)

    def test_table_paragraphs_are_detected_and_preserved(self):
        make_resume(self.settings.base_resume, table=True)
        agent = self.agent()
        index = next(
            i
            for i, paragraph in enumerate(agent.document.paragraphs)
            if paragraph.text.startswith("Maintained")
        )
        self.assertIn(index, agent.targets)
        output = self.settings.output_root / "tailored_resumes/table.docx"
        agent.apply(
            [
                {
                    "paragraph_id": index,
                    "replacement": "Maintained reliable Python service tests supporting delivery across collaborative product teams.",
                }
            ],
            output,
        )
        self.assertIn("collaborative", ResumeDocument(output).paragraphs[index].text)
        self.assertEqual(len(Document(output).tables), 1)

    def test_package_parts_and_source_remain_unchanged(self):
        agent = self.agent()
        index = self.target(agent)
        source_hash = digest(agent.base.read_bytes())
        output = self.settings.output_root / "tailored_resumes/copy.docx"
        agent.apply(
            [
                {
                    "paragraph_id": index,
                    "replacement": "Software professional delivering reliable Python backend services and automated testing for collaborative product teams.",
                }
            ],
            output,
        )
        self.assertEqual(source_hash, digest(agent.base.read_bytes()))
        with ZipFile(agent.base) as source, ZipFile(output) as result:
            self.assertEqual(source.namelist(), result.namelist())
            for name in source.namelist():
                if name != "word/document.xml":
                    self.assertEqual(source.read(name), result.read(name))
        for i, paragraph in enumerate(ResumeDocument(output).paragraphs):
            if i != index:
                self.assertEqual(paragraph.text, agent.document.paragraphs[i].text)

    def test_protected_target_and_new_numeric_claim_rejected(self):
        agent = self.agent()
        with self.assertRaises(ResumeError):
            agent.validate_edits([{"paragraph_id": 0, "replacement": "Other identity"}])
        with self.assertRaisesRegex(ResumeError, "unsupported number"):
            agent.validate_edits(
                [
                    {
                        "paragraph_id": self.target(agent),
                        "replacement": "Delivered 99% availability using Python.",
                    }
                ]
            )

    def test_advisory_keeps_unsupported_number_as_warning(self):
        self.raw["resume"]["factual_review_mode"] = "advisory"
        agent = self.agent()
        audit = agent.validate_edits(
            [
                {
                    "paragraph_id": self.target(agent),
                    "replacement": "Delivered 99% availability using Python.",
                }
            ]
        )
        self.assertTrue(audit[0]["validation_warnings"])

    def test_skills_prose_and_label_changes_rejected(self):
        agent = self.agent()
        index = self.target(agent, "skills")
        for text in [
            "Different: Python",
            "Programming: Python, responsible for building scalable systems and supporting product delivery",
        ]:
            with self.assertRaises(ResumeError):
                agent.validate_edits([{"paragraph_id": index, "replacement": text}])

    def test_uncategorized_skill_lists_supported(self):
        doc = Document(self.settings.base_resume)
        doc.paragraphs[5].text = "Python, SQL"
        doc.save(self.settings.base_resume)
        agent = self.agent()
        index = self.target(agent, "skills")
        self.assertTrue(
            agent.validate_edits([{"paragraph_id": index, "replacement": "SQL, Python, testing"}])
        )

    def test_duplicate_bullets_rejected_in_all_modes(self):
        for mode in ["strict", "advisory"]:
            self.raw["resume"]["factual_review_mode"] = mode
            agent = self.agent()
            ids = [
                index
                for index, entry in agent.targets.items()
                if entry["section"] == "experience_bullets"
            ]
            edits = [
                {
                    "paragraph_id": index,
                    "replacement": "Built reliable Python services and automated checks supporting collaborative product delivery.",
                }
                for index in ids
            ]
            self.assertTrue(agent.quality_issues(edits))

    def test_custom_heading_alias_and_unknown_layout_error(self):
        doc = Document()
        doc.add_heading("My Craft", 1)
        doc.add_paragraph("Python, SQL")
        doc.save(self.settings.base_resume)
        with self.assertRaisesRegex(ResumeError, "section_aliases"):
            self.agent()
        self.raw["resume"]["section_aliases"] = {"skills": ["My Craft"]}
        self.assertEqual(len(self.agent().targets), 1)

    def test_complex_word_paragraph_is_protected(self):
        doc = Document()
        paragraph = doc.add_paragraph("Python\tSQL")
        with self.assertRaises(ResumeError):
            replace_text_nodes(paragraph, "Other text")

    def test_runs_formatting_and_parenthesized_skill_lists(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("Reliable ").bold = True
        paragraph.add_run("Python services").italic = True
        replace_text_nodes(paragraph, "Reliable Python backend services")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertTrue(paragraph.runs[1].italic)
        self.assertEqual(
            skill_items("Python, Tool (A, B); Testing"), ["Python", "Tool (A, B)", "Testing"]
        )

    def test_full_tailoring_strict_and_advisory_review(self):
        for mode, supported in [("strict", True), ("advisory", False)]:
            self.raw["resume"]["factual_review_mode"] = mode
            agent = self.agent()
            index = self.target(agent)

            class Model:
                def generate(self, role, instructions, payload, schema):
                    if role == "resume_tailorer":
                        return {
                            "edits": [
                                {
                                    "paragraph_id": index,
                                    "replacement": "Software professional delivering reliable Python backend services and testing for collaborative product teams.",
                                }
                            ]
                        }
                    return {
                        "decisions": [
                            {
                                "paragraph_id": row["paragraph_id"],
                                "supported": supported,
                                "reason": "Test factual assessment",
                            }
                            for row in payload["proposed_edits"]
                        ]
                    }

            agent.llm = Model()
            result = agent.tailor(job())
            self.assertTrue(Path(result["resume_path"]).is_file())
            self.assertEqual(result["manual_fact_review_required"], mode == "advisory")
            self.assertFalse(result["pagination_verified"])

    def test_strict_reviewer_rejection_blocks_output(self):
        agent = self.agent()
        index = self.target(agent)

        class Model:
            def generate(self, role, instructions, payload, schema):
                if role == "resume_tailorer":
                    return {
                        "edits": [
                            {
                                "paragraph_id": index,
                                "replacement": "Software professional delivering Python systems for product teams.",
                            }
                        ]
                    }
                return {
                    "decisions": [
                        {"paragraph_id": index, "supported": False, "reason": "Unsupported claim"}
                    ]
                }

        agent.llm = Model()
        with self.assertRaisesRegex(ResumeError, "rejected all edits"):
            agent.tailor(job())
        self.assertFalse(list(self.settings.output_root.rglob("*.docx")))


class UnknownHeadingTests(ProjectCase):
    def test_unknown_heading_after_skills_remains_protected(self):
        doc = Document()
        doc.add_heading("Skills", 1)
        doc.add_paragraph("Python, SQL")
        doc.add_heading("Additional Background", 1)
        doc.add_paragraph(
            "Completed a computing degree at an example university and professional certification training."
        )
        doc.save(self.settings.base_resume)
        agent = ResumeAgent(self.load(), None)
        target_text = [agent.document.paragraphs[index].text for index in agent.targets]
        self.assertEqual(target_text, ["Python, SQL"])

    def test_bold_skills_category_is_still_editable(self):
        doc = Document()
        doc.add_heading("Skills", 1)
        category = doc.add_paragraph()
        category.add_run("Programming: Python, SQL").bold = True
        doc.save(self.settings.base_resume)
        agent = ResumeAgent(self.load(), None)
        self.assertEqual(len(agent.targets), 1)
