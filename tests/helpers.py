"""Synthetic fixtures only; tests never read a real user's resume or settings."""

from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path
from typing import Any

from docx import Document

from config_loader import PROJECT_ROOT, Settings, load_settings
from jobfinder.common import utcnow
from jobfinder.models import Job


def make_resume(path: Path, *, table: bool = False, shifted: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("Example Candidate")
    doc.add_paragraph("candidate@example.com")
    if shifted:
        doc.add_paragraph("")
        doc.add_paragraph("")
    doc.add_heading("Professional Summary", 1)
    doc.add_paragraph(
        "Software professional building reliable backend services with Python and testing practices for collaborative product teams across several delivery environments."
    )
    doc.add_heading("Skills", 1)
    doc.add_paragraph("Programming: Python, SQL")
    doc.add_heading("Professional Experience", 1)
    doc.add_paragraph("Example Employer | Software Developer | 2021–Present")
    doc.add_paragraph(
        "Built Python services supporting reliable API delivery and automated checks for internal product teams.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Improved SQL queries for reporting workflows and supported production troubleshooting with colleagues.",
        style="List Bullet",
    )
    if table:
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[
            0
        ].text = "Maintained Python service tests for reliable delivery across product teams."
        cell.paragraphs[0].style = "List Bullet"
    doc.add_heading("Education", 1)
    doc.add_paragraph("Example University | Computing Degree | 2020")
    doc.save(path)


def job(index: int = 0, **kwargs: Any) -> Job:
    return Job(
        job_title="Software Engineer",
        company=f"Example Company {index}",
        location="London, UK",
        job_description="Build Python services and reliable APIs for product teams in London, UK. Collaborate on automated testing and SQL reporting. "
        * 2,
        application_url=f"https://example.com/jobs/{index}",
        posted_at=utcnow().isoformat(),
        **kwargs,
    )


def analysis(**overrides: Any) -> dict[str, Any]:
    return {
        "target_role_match": True,
        "location_eligible": True,
        "location_evidence": "London, UK",
        "employment_type": "full_time",
        "work_arrangement": "hybrid",
        "sponsorship_status": "unknown",
        "sponsorship_evidence": [],
        "required_experience": "Not stated",
        "required_experience_min": None,
        "required_experience_max": None,
        "experience_evidence": "",
        "experience_level": "unknown",
        "experience_level_match": True,
        "qualifications": [],
        **overrides,
    }


class ProjectCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        shutil.copytree(PROJECT_ROOT / "skills", self.root / "skills")
        self.raw = {
            "search": {
                "job_titles": ["Software Engineer"],
                "country": "UK",
                "locations": ["London"],
                "verify_employer": False,
                "batch_size": 2,
            },
            "llm": {
                "models": {
                    role: "example-local:latest"
                    for role in [
                        "job_parser",
                        "keyword_extractor",
                        "resume_tailorer",
                        "resume_reviewer",
                    ]
                }
            },
            "resume": {"base_path": "templates/base.docx"},
            "email": {"enabled": False},
        }
        make_resume(self.root / "templates/base.docx")
        self.settings = self.load()

    def load(self) -> Settings:
        path = self.root / "config.json"
        path.write_text(json.dumps(self.raw), encoding="utf-8")
        return load_settings(path, project_root=self.root)
